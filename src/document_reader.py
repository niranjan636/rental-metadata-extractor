
from pathlib import Path
from docx import Document


def read_docx(file_path: str | Path) -> str:
    """Extract paragraphs and table text from a DOCX file."""
    document = Document(str(file_path))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)
